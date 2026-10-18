# -*- coding: utf-8 -*-
"""기록 저장.

- 원본 기록은 '기록데이터' 폴더의 파일(json)에 쌓입니다. (엑셀이 열려 있어도 기록이 사라지지 않게)
- 기록할 때마다 '랭챌기록.xlsx'를 새로 만듭니다. 엑셀이 열려 있으면 저장을 건너뛰고 다음에 다시 만듭니다.
"""
import os, json, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

APP_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT_DIR = os.path.dirname(APP_DIR)
DATA_DIR = os.path.join(ROOT_DIR, '기록데이터')
XLSX = os.path.join(ROOT_DIR, '랭챌기록.xlsx')
os.makedirs(DATA_DIR, exist_ok=True)

MY_DECK = os.path.join(DATA_DIR, 'my_deck.json')
PENDING = os.path.join(DATA_DIR, 'pending_analysis.json')
MATCHES = os.path.join(DATA_DIR, 'matches.jsonl')
PLAYERS = os.path.join(DATA_DIR, 'players.jsonl')
ANALYSES = os.path.join(DATA_DIR, 'analyses.jsonl')


def now():
    return datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')


def save_json(path, obj):
    tmp = path + '.tmp'
    with open(tmp, 'w', encoding='utf-8') as f:
        json.dump(obj, f, ensure_ascii=False, indent=1)
    os.replace(tmp, path)


def load_json(path, default=None):
    if not os.path.isfile(path):
        return default
    with open(path, encoding='utf-8') as f:
        return json.load(f)


def append_jsonl(path, rows):
    with open(path, 'a', encoding='utf-8') as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + '\n')


def read_jsonl(path):
    if not os.path.isfile(path):
        return []
    out = []
    with open(path, encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if line:
                try:
                    out.append(json.loads(line))
                except json.JSONDecodeError:
                    pass
    return out


# ============================================================ 엑셀
MATCH_COLS = ['기록시각', '분석시각', '줄', '상대닉네임', '상대POINT', '상대OVR(로비)', '상대타자OVR',
              '상대선발', '상대선발손', '상대선발순번', '내선발', '내선발손', '내선발순번', '추천선발순번', '추천따름',
              '내타선점수', '내선발점수', '상대타선점수', '상대선발점수', '내합계', '상대합계', '점수차', '예상승률',
              '내팀OVR', '상대팀OVR', '결과', '스코어', '내득점', '상대득점', '분석제외']
PLAYER_COLS = ['기록시각', '분석시각', '줄', '팀', '역할', '타순/순번', '이름', '포지션', 'OVR', '카드OVR', '스탯평균OVR', '별', '카드종류',
               '투', '타', '파워/구속', '정확/변화', '선구/구위', '인내/제구', '주루/지구력', '수비', '발사각',
               '잠재1', '잠재2', '잠재3', '잠재4', '잠재5', '각잠', '스킬1', 'Lv1', '점수1', '스킬2', 'Lv2', '점수2',
               '스킬3', 'Lv3', '점수3', '능력치점수', '잠재점수', '기타보정', '스킬합', '매치업보정점수']
DECK_COLS = ['구분', '포지션/선발', '타순', '이름', 'OVR', '별', '카드종류', '투', '타',
             '파워/구속', '정확/변화', '선구/구위', '인내/제구', '주루/지구력', '수비', '발사각',
             '잠재1', '잠재2', '잠재3', '잠재4', '잠재5', '각잠', '스킬1', 'Lv1', '스킬2', 'Lv2', '스킬3', 'Lv3',
             '평균조건 보정점수', '읽은시각']
ANALYSIS_COLS = ['분석시각', '줄', '상대닉네임', '상대선발', '상대선발손', '추천 내선발', '점수차합 기준 추천',
                 '내1선발 점수차', '내2선발 점수차', '내3선발 점수차', '내4선발 점수차', '내5선발 점수차'] + \
                [f'내{i}선발 예상승률' for i in range(1, 6)] + \
                [f'내{i}선발 투수점수' for i in range(1, 6)] + [f'상대타선(vs 내{i}선발)' for i in range(1, 6)] + \
                ['내 타선점수', '상대 선발점수']


def player_row_values(p):
    pit = p.get('pitcher')
    g = (p.get('grades') or []) + [''] * 6
    sk = (p.get('skills') or []) + [('', None)] * 3
    if pit:
        st = [p.get('구속'), p.get('변화'), p.get('구위'), p.get('제구'), p.get('지구력'), p.get('수비'), None]
    else:
        st = [p.get('파워'), p.get('정확'), p.get('선구'), p.get('인내'), p.get('주루'), p.get('수비'), p.get('발사각')]
    return st, g[:6], sk[:3]


def _sheet(wb, title, cols, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(cols)
    for r in rows:
        ws.append([r.get(c) for c in cols])
    head = PatternFill('solid', fgColor='1F3864')
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = head
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.freeze_panes = 'A2'
    for i, col in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(col, max(8, min(16, len(col) * 2 + 2)))
    return ws


def deck_rows(deck):
    rows = []
    for p in (deck or {}).get('batters', []) + (deck or {}).get('pitchers', []):
        st, g, sk = player_row_values(p)
        pit = p.get('pitcher')
        rows.append({
            '구분': '선발투수' if pit else '타자',
            '포지션/선발': f"{p.get('rotation')}선발" if pit else p.get('pos'),
            '타순': None if pit else p.get('slot'),
            '이름': p.get('name'), 'OVR': p.get('ovr'), '별': p.get('stars'), '카드종류': p.get('card_type'),
            '투': p.get('throw'), '타': p.get('bat'),
            '파워/구속': st[0], '정확/변화': st[1], '선구/구위': st[2], '인내/제구': st[3], '주루/지구력': st[4],
            '수비': st[5], '발사각': st[6],
            '잠재1': g[0], '잠재2': g[1], '잠재3': g[2], '잠재4': g[3], '잠재5': g[4], '각잠': g[5],
            '스킬1': sk[0][0], 'Lv1': sk[0][1], '스킬2': sk[1][0], 'Lv2': sk[1][1], '스킬3': sk[2][0], 'Lv3': sk[2][1],
            '평균조건 보정점수': round(p.get('avg_score') or 0, 2), '읽은시각': deck.get('time')})
    return rows


def rebuild_excel():
    """모든 기록으로 엑셀을 새로 만듦. 성공하면 None, 실패하면 오류 메시지"""
    wb = Workbook()
    wb.remove(wb.active)
    info = wb.create_sheet('안내')
    for line in ['랭챌 분석기 기록 파일',
                 '',
                 '매치기록: 경기 한 줄 = 한 경기. 점수는 모두 매치업(이번 상대 조건) 기준입니다.',
                 '선수기록: 경기마다 양 팀 선수 10명씩의 원본값과 점수.',
                 '분석기록: 상대 분석 때 계산한 5x5 표 (내 선발 x 상대). 점수차와 그 구성(내 선발 점수, 상대 타선 점수 등).',
                 '내 덱: 마지막으로 [① 내 덱 읽기] 한 결과.',
                 '',
                 '※ 이 파일은 프로그램이 매번 새로 만듭니다. 직접 고친 내용은 저장한 뒤 프로그램의 [엑셀 수정 반영]을 눌러야 원본에 들어가요.',
                 '   (매치기록: 결과/스코어/이름 등, 선수기록: 이름/스탯/스킬/잠재/좌우 등, 내 덱: 모든 값. 점수 칸은 자동으로 다시 계산돼요)',
                 '※ 원본은 기록데이터 폴더에 있습니다. 이 파일을 지워도 다음 기록 때 다시 만들어집니다.']:
        info.append([line])
    info.column_dimensions['A'].width = 90
    _sheet(wb, '매치기록', MATCH_COLS, read_jsonl(MATCHES))
    _sheet(wb, '선수기록', PLAYER_COLS, read_jsonl(PLAYERS))
    _sheet(wb, '분석기록', ANALYSIS_COLS, read_jsonl(ANALYSES))
    _sheet(wb, '내 덱', DECK_COLS, deck_rows(load_json(MY_DECK)))
    try:
        wb.save(XLSX)
        return None
    except PermissionError:
        return '랭챌기록.xlsx가 엑셀에서 열려 있어서 저장하지 못했어요. 엑셀을 닫으면 다음 기록 때 함께 저장돼요. (기록 자체는 안전하게 남아 있어요)'
